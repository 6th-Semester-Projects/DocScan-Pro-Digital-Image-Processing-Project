import docx
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

# Theme Colors
COLOR_BLUE_DARK = RGBColor(0x1A, 0x5C, 0xB5)
COLOR_BLUE_LIGHT = RGBColor(0xE8, 0xF0, 0xFE)
COLOR_GRAY = RGBColor(0x55, 0x55, 0x55)
COLOR_GRAY_BG = RGBColor(0xF0, 0xF4, 0xF8)

def setup_document(doc):
    """Setup document margins, fonts, and styles according to prompt."""
    # 5. MARGINS: Top/Bottom 2.54cm, Left 3.18cm, Right 2.54cm
    sections = doc.sections
    for section in sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.18)
        section.right_margin = Cm(2.54)
        
    # 2. FONTS: Times New Roman for body text (12pt), 1.5 line spacing, 6pt after
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    pf = style.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # Create heading styles
    _setup_heading_style(doc, 'Heading 1', 16, True)
    _setup_heading_style(doc, 'Heading 2', 14, False)
    _setup_heading_style(doc, 'Heading 3', 12, False)

def _setup_heading_style(doc, style_name, size, page_break_before):
    style = doc.styles[style_name]
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(size)
    font.color.rgb = COLOR_BLUE_DARK
    font.bold = True
    pf = style.paragraph_format
    pf.space_before = Pt(12)
    pf.space_after = Pt(6)
    pf.keep_with_next = True
    if page_break_before:
        pf.page_break_before = True

def add_blue_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    return h

def add_para(doc, text, bold=False, italic=False, align="justify"):
    p = doc.add_paragraph()
    if align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "left":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        
    run = p.add_run(text)
    if bold: run.bold = True
    if italic: run.italic = True
    return p

def set_cell_bg(cell, fill_color):
    """Set background color of a cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)

def _set_cell_borders(cell):
    """Add standard borders to a cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for border_name in ['top', 'left', 'bottom', 'right']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '1A5CB5')
        tcBorders.append(border)
    tcPr.append(tcBorders)

def add_blue_table(doc, headers, rows_data, col_widths=None):
    """Create a styled table according to rules."""
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    
    # Header Row
    hdr_cells = table.rows[0].cells
    for i, hdr in enumerate(headers):
        hdr_cells[i].text = str(hdr)
        set_cell_bg(hdr_cells[i], "1A5CB5")
        
        # Format text
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            
    # Data Rows
    for r_idx, row_data in enumerate(rows_data):
        row_cells = table.add_row().cells
        bg_color = "E8F0FE" if r_idx % 2 != 0 else "FFFFFF"
        for i, val in enumerate(row_data):
            row_cells[i].text = str(val)
            set_cell_bg(row_cells[i], bg_color)
            
            p = row_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(10)
                
    if col_widths:
        for i, width in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = width
                
    doc.add_paragraph() # Empty paragraph after table
    return table

def add_image_with_caption(doc, image_path, caption_text, width_inches=5.5):
    import os
    if not os.path.exists(image_path):
        print(f"Warning: Image not found at {image_path}")
        return
        
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(image_path, width=Inches(width_inches))
    
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c_run = cp.add_run(caption_text)
    c_run.italic = True
    c_run.font.name = 'Times New Roman'
    c_run.font.size = Pt(10)
    c_run.font.color.rgb = COLOR_GRAY
    
    doc.add_paragraph() # Empty paragraph after image

def add_section_divider(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━")
    run.font.color.rgb = COLOR_BLUE_LIGHT
    run.font.size = Pt(10)

def add_info_box(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    cell = table.cell(0, 0)
    set_cell_bg(cell, "E8F0FE")
    _set_cell_borders(cell)
    
    # Title
    p1 = cell.paragraphs[0]
    r1 = p1.add_run(f"ℹ {title}\n")
    r1.font.name = 'Calibri'
    r1.font.size = Pt(11)
    r1.font.bold = True
    r1.font.color.rgb = COLOR_BLUE_DARK
    
    # Body
    r2 = p1.add_run(body)
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(10)

    doc.add_paragraph()

def add_code_block(doc, code, caption=None):
    if caption:
        p = doc.add_paragraph()
        r = p.add_run(caption)
        r.font.name = 'Calibri'
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = COLOR_BLUE_DARK
        
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    cell = table.cell(0, 0)
    set_cell_bg(cell, "F0F4F8")
    _set_cell_borders(cell)
    
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(code)
    r.font.name = 'Consolas'
    r.font.size = Pt(8)
    r.font.color.rgb = COLOR_BLUE_DARK
    
    doc.add_paragraph()

def add_header_footer(doc, project_title):
    for section in doc.sections:
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        hr = hp.add_run(project_title)
        hr.italic = True
        hr.font.name = 'Calibri'
        hr.font.size = Pt(9)
        hr.font.color.rgb = COLOR_BLUE_DARK
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Auto page number magic in python-docx
        run = fp.add_run()
        fldChar1 = OxmlElement('w:fldChar')
        fldChar1.set(qn('w:fldCharType'), 'begin')
        instrText = OxmlElement('w:instrText')
        instrText.set(qn('xml:space'), 'preserve')
        instrText.text = "PAGE"
        fldChar2 = OxmlElement('w:fldChar')
        fldChar2.set(qn('w:fldCharType'), 'end')
        
        run._r.append(fldChar1)
        run._r.append(instrText)
        run._r.append(fldChar2)

def add_toc(doc):
    p = doc.add_paragraph()
    r = p.add_run()
    fldChar = OxmlElement('w:fldChar')
    fldChar.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    
    r._r.append(fldChar)
    r._r.append(instrText)
    r._r.append(fldChar2)
    r._r.append(fldChar3)
